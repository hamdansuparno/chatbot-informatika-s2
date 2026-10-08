#!/usr/bin/env python3
"""
office_tool.py - Baca / tulis / terjemahkan file Excel, Word, PDF, PowerPoint.
Read / write / translate Excel, Word, PDF, PowerPoint files.
读取 / 写入 / 翻译 Excel、Word、PDF、PowerPoint 文件。

Instalasi / Install:
    pip install openpyxl python-docx pypdf reportlab python-pptx deep-translator

Contoh / Examples:
    python office_tool.py read   data.xlsx
    python office_tool.py write  hasil.docx --text "Halo dunia" --text "Baris kedua"
    python office_tool.py write  hasil.pptx --input isi.txt      # slide dipisah baris kosong
    python office_tool.py translate laporan.docx laporan_en.docx --to en
    python office_tool.py translate data.xlsx data_zh.xlsx --to zh
Bahasa / Languages: id (Indonesia), en (English), zh (中文)
"""
import argparse
import sys
from pathlib import Path
from xml.sax.saxutils import escape

LANGS = {"id": "id", "en": "en", "zh": "zh-CN"}
_cache = {}


# ------------------------------------------------------------------ translate
def tr(text, target):
    """Terjemahkan teks ke bahasa target (butuh internet)."""
    t = text.strip()
    if not t:
        return text
    key = (t, target)
    if key not in _cache:
        from deep_translator import GoogleTranslator
        g = GoogleTranslator(source="auto", target=LANGS[target])
        parts = [t[i:i + 4500] for i in range(0, len(t), 4500)]
        _cache[key] = "".join(g.translate(p) or p for p in parts)
    return _cache[key]


def _retext(par, lang):
    """Ganti teks paragraf (docx/pptx), format run pertama dipertahankan."""
    if par.runs and par.text.strip():
        par.runs[0].text = tr(par.text, lang)
        for r in par.runs[1:]:
            r.text = ""


# ---------------------------------------------------------------------- Excel
def read_xlsx(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)
    return {ws.title: [list(r) for r in ws.iter_rows(values_only=True)] for ws in wb}


def write_xlsx(path, lines):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    for line in lines:
        ws.append(line.split(","))  # pisahkan kolom dengan koma
    wb.save(path)


# ----------------------------------------------------------------------- Word
def read_docx(path):
    from docx import Document
    d = Document(path)
    out = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            out.append(" | ".join(c.text for c in row.cells))
    return out


def write_docx(path, lines):
    from docx import Document
    d = Document()
    for line in lines:
        d.add_paragraph(line)
    d.save(path)


# ------------------------------------------------------------------------ PDF
def read_pdf(path):
    from pypdf import PdfReader
    return [(p.extract_text() or "") for p in PdfReader(path).pages]


def write_pdf(path, lines):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
    # Font CJK bawaan reportlab: mendukung Latin + Mandarin
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    style = ParagraphStyle("body", fontName="STSong-Light", fontSize=11,
                           leading=16, wordWrap="CJK")
    story = []
    for line in lines:
        for sub in str(line).split("\n"):
            story.append(Paragraph(escape(sub) or "&nbsp;", style))
        story.append(Spacer(1, 4))
    SimpleDocTemplate(str(path), pagesize=A4).build(story)


# ----------------------------------------------------------------- PowerPoint
def read_pptx(path):
    from pptx import Presentation
    out = []
    for i, slide in enumerate(Presentation(path).slides, 1):
        out.append(f"[Slide {i}]")
        for sh in slide.shapes:
            if sh.has_text_frame:
                out += [p.text for p in sh.text_frame.paragraphs if p.text]
    return out


def write_pptx(path, lines):
    from pptx import Presentation
    prs = Presentation()
    blocks, cur = [], []
    for line in lines + [""]:
        if line.strip():
            cur.append(line)
        elif cur:
            blocks.append(cur)
            cur = []
    for b in blocks:  # baris pertama = judul, sisanya = isi
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = b[0]
        s.placeholders[1].text = "\n".join(b[1:])
    prs.save(path)


# ------------------------------------------------------- translate whole file
def translate_file(src, dst, lang):
    ext = src.suffix.lower()
    if ext == ".xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(src)
        for ws in wb:
            for row in ws.iter_rows():
                for c in row:
                    if isinstance(c.value, str) and not c.value.startswith("="):
                        c.value = tr(c.value, lang)
        wb.save(dst)
    elif ext == ".docx":
        from docx import Document
        d = Document(src)
        pars = list(d.paragraphs)
        for t in d.tables:
            for row in t.rows:
                for c in row.cells:
                    pars += c.paragraphs
        for p in pars:
            _retext(p, lang)
        d.save(dst)
    elif ext == ".pptx":
        from pptx import Presentation
        prs = Presentation(src)
        for slide in prs.slides:
            for sh in slide.shapes:
                if sh.has_text_frame:
                    for p in sh.text_frame.paragraphs:
                        _retext(p, lang)
        prs.save(dst)
    elif ext == ".pdf":
        write_pdf(dst, [tr(t, lang) for t in read_pdf(src) if t.strip()])
    else:
        sys.exit(f"Format tidak didukung: {ext}")


# ------------------------------------------------------------------------ CLI
READERS = {".xlsx": read_xlsx, ".docx": read_docx, ".pdf": read_pdf, ".pptx": read_pptx}
WRITERS = {".xlsx": write_xlsx, ".docx": write_docx, ".pdf": write_pdf, ".pptx": write_pptx}


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    r = sub.add_parser("read")
    r.add_argument("file")
    w = sub.add_parser("write")
    w.add_argument("file")
    w.add_argument("--text", action="append", default=[], help="satu baris/paragraf")
    w.add_argument("--input", help="file .txt sumber isi")
    t = sub.add_parser("translate")
    t.add_argument("src")
    t.add_argument("dst")
    t.add_argument("--to", choices=LANGS, required=True)
    a = ap.parse_args()

    if a.cmd == "read":
        p = Path(a.file)
        data = READERS[p.suffix.lower()](p)
        if isinstance(data, dict):
            for name, rows in data.items():
                print(f"## {name}")
                for row in rows:
                    print("\t".join("" if v is None else str(v) for v in row))
        else:
            print("\n".join(data))
    elif a.cmd == "write":
        p = Path(a.file)
        lines = list(a.text)
        if a.input:
            lines += Path(a.input).read_text(encoding="utf-8").splitlines()
        WRITERS[p.suffix.lower()](p, lines)
        print("Tersimpan:", p)
    else:
        translate_file(Path(a.src), Path(a.dst), a.to)
        print("Tersimpan:", a.dst)


if __name__ == "__main__":
    main()
